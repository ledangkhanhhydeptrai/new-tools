# ============================================================
# merge_google_maps_to_excel.py
#
# MỤC ĐÍCH:
#   1. Đọc file Excel CRM chính
#   2. Đọc file Excel đã check Google Maps
#   3. Match theo tên khách sạn
#   4. Giữ nguyên Address
#   5. Tạo cột "Google Maps" ngay bên phải Address
#   6. Nếu đã có link Maps rồi -> giữ nguyên
#   7. Nếu chưa có -> lấy Maps từ file checked và gắn vào
#   8. Xuất file Excel mới
#
# CÀI:
#   pip install pandas openpyxl
# ============================================================

import os
import re
import sys
import unicodedata
from typing import Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


# ============================================================
# CONFIG
# ============================================================

MAIN_FILE = "hotels_export(1).xlsx"

MAPS_FILE = "hotels_export_google_maps_checked(2).xlsx"

OUTPUT_FILE = "hotels_export_with_google_maps.xlsx"


# ============================================================
# GOOGLE MAPS URL
# ============================================================

GOOGLE_MAPS_PATTERNS = (
    "google.com/maps",
    "maps.google.",
    "maps.app.goo.gl",
    "goo.gl/maps",
)


def is_google_maps_url(value) -> bool:
    """
    Kiểm tra value có phải link Google Maps hay không.
    """

    if value is None:
        return False

    try:
        if pd.isna(value):
            return False
    except Exception:
        pass

    text = str(value).strip().lower()

    if not text:
        return False

    return any(pattern in text for pattern in GOOGLE_MAPS_PATTERNS)


# ============================================================
# NORMALIZE
# ============================================================


def safe_text(value) -> str:
    """
    Chuyển mọi value thành string an toàn.
    """

    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    return str(value).strip()


def normalize_text(value) -> str:
    """
    Normalize text để match tên khách sạn.

    Ví dụ:
        "Khách Sạn Dakruco"
        "khach san dakruco"

    -> đều gần như cùng key
    """

    text = safe_text(value)

    if not text:
        return ""

    text = unicodedata.normalize(
        "NFKD",
        text,
    )

    text = "".join(char for char in text if not unicodedata.combining(char))

    text = text.lower()

    text = re.sub(
        r"[^\w\s]",
        " ",
        text,
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


# ============================================================
# FIND HEADER ROW
# ============================================================


def find_header_row(
    file_path: str,
    required_columns=("Tên khách sạn", "Address"),
    max_rows: int = 20,
) -> int:
    """
    Tự tìm dòng header trong file CRM.

    File của bạn có dạng:

        dòng 1: Title / Email / Phone / Address ...
        dòng 2: trống
        dòng 3: Tên khách sạn / Email / Phone / Stars ... Address

    Hàm này sẽ tự tìm dòng 3.

    Return:
        index dùng cho pandas header=
        tức Excel row 3 -> pandas index 2
    """

    preview = pd.read_excel(
        file_path,
        header=None,
        nrows=max_rows,
    )

    required_normalized = {normalize_text(column) for column in required_columns}

    for row_index, row in preview.iterrows():
        current = {normalize_text(value) for value in row.tolist() if safe_text(value)}

        if required_normalized.issubset(current):
            return int(row_index)

    raise ValueError(f"Không tìm thấy dòng header chứa {required_columns}")


# ============================================================
# FIND COLUMN
# ============================================================


def find_column(
    df: pd.DataFrame,
    candidates,
) -> Optional[str]:
    """
    Tìm tên column không phân biệt hoa thường / dấu.
    """

    normalized_columns = {normalize_text(column): column for column in df.columns}

    for candidate in candidates:
        key = normalize_text(candidate)

        if key in normalized_columns:
            return normalized_columns[key]

    return None


# ============================================================
# LOAD MAIN EXCEL
# ============================================================


def load_main_excel(
    file_path: str,
):
    """
    Load file CRM chính.
    """

    header_row = find_header_row(
        file_path,
        required_columns=(
            "Tên khách sạn",
            "Address",
        ),
    )

    df = pd.read_excel(
        file_path,
        header=header_row,
        dtype=object,
    )

    return df, header_row


# ============================================================
# LOAD GOOGLE MAPS FILE
# ============================================================


def load_maps_excel(
    file_path: str,
):
    """
    Load file kết quả Google Maps.
    """

    df = pd.read_excel(
        file_path,
        dtype=object,
    )

    title_col = find_column(
        df,
        (
            "Title",
            "Tên khách sạn",
            "Hotel",
            "Hotel Name",
            "Name",
        ),
    )

    if not title_col:
        raise ValueError("Không tìm thấy cột Title trong file Google Maps.")

    google_maps_col = find_column(
        df,
        (
            "google_maps_url",
            "Google Maps",
            "Google Maps URL",
            "maps_url",
            "maps_check_url",
        ),
    )

    url_col = find_column(
        df,
        (
            "URL",
            "Url",
            "url",
        ),
    )

    return (
        df,
        title_col,
        google_maps_col,
        url_col,
    )


# ============================================================
# BUILD MAP LOOKUP
# ============================================================


def build_google_maps_lookup(
    maps_df: pd.DataFrame,
    title_col: str,
    google_maps_col: Optional[str],
    url_col: Optional[str],
):
    """
    Tạo dictionary:

        normalized hotel name
            ->
        google maps url
    """

    lookup = {}

    for _, row in maps_df.iterrows():
        title = safe_text(row.get(title_col))

        if not title:
            continue

        maps_url = ""

        # Ưu tiên google_maps_url
        if google_maps_col:
            value = row.get(google_maps_col)

            if is_google_maps_url(value):
                maps_url = safe_text(value)

        # fallback URL
        if not maps_url and url_col:
            value = row.get(url_col)

            if is_google_maps_url(value):
                maps_url = safe_text(value)

        if not maps_url:
            continue

        key = normalize_text(title)

        if not key:
            continue

        # Không overwrite nếu đã có URL
        if key not in lookup:
            lookup[key] = maps_url

    return lookup


# ============================================================
# INSERT GOOGLE MAPS COLUMN
# ============================================================


def merge_google_maps(
    main_df: pd.DataFrame,
    maps_lookup: dict,
):
    """
    Gắn Google Maps vào cột ngay bên phải Address.
    """

    title_col = find_column(
        main_df,
        (
            "Tên khách sạn",
            "Title",
            "Hotel",
            "Hotel Name",
            "Name",
        ),
    )

    address_col = find_column(
        main_df,
        (
            "Address",
            "Địa chỉ",
            "Dia chi",
        ),
    )

    if not title_col:
        raise ValueError("Không tìm thấy cột Tên khách sạn.")

    if not address_col:
        raise ValueError("Không tìm thấy cột Address.")

    # ========================================================
    # Nếu đã có Google Maps column
    # ========================================================

    existing_maps_col = find_column(
        main_df,
        (
            "Google Maps",
            "Google Maps URL",
            "google_maps_url",
            "maps_url",
        ),
    )

    # ========================================================
    # Nếu chưa có thì tạo ngay sau Address
    # ========================================================

    if not existing_maps_col:
        address_index = main_df.columns.get_loc(address_col)

        main_df.insert(
            address_index + 1,
            "Google Maps",
            "",
        )

        maps_col = "Google Maps"

    else:
        maps_col = existing_maps_col

    # ========================================================
    # PROCESS
    # ========================================================

    matched = 0
    already_has = 0
    not_found = 0
    empty_title = 0

    for index, row in main_df.iterrows():
        title = safe_text(row.get(title_col))

        if not title:
            empty_title += 1
            continue

        # ------------------------------------
        # Nếu Google Maps column đã có link
        # ------------------------------------

        current_maps = row.get(maps_col)

        if is_google_maps_url(current_maps):
            already_has += 1
            continue

        # ------------------------------------
        # Một số file cũ có thể đã nhét Maps
        # vào Address
        # ------------------------------------

        address = safe_text(row.get(address_col))

        if is_google_maps_url(address):
            already_has += 1
            continue

        # ------------------------------------
        # Match title
        # ------------------------------------

        key = normalize_text(title)

        maps_url = maps_lookup.get(
            key,
            "",
        )

        if not maps_url:
            not_found += 1
            continue

        main_df.at[
            index,
            maps_col,
        ] = maps_url

        matched += 1

    return {
        "df": main_df,
        "title_col": title_col,
        "address_col": address_col,
        "maps_col": maps_col,
        "matched": matched,
        "already_has": already_has,
        "not_found": not_found,
        "empty_title": empty_title,
    }


# ============================================================
# SAVE WHILE KEEPING ORIGINAL LAYOUT
# ============================================================


def save_excel(
    original_file: str,
    output_file: str,
    df: pd.DataFrame,
    header_row: int,
    address_col: str,
    maps_col: str,
):
    """
    Ghi dữ liệu trở lại Excel.

    Giữ các dòng phía trên header giống file gốc.
    """

    # ========================================================
    # Copy file trước
    # ========================================================

    import shutil

    shutil.copy2(
        original_file,
        output_file,
    )

    # ========================================================
    # Load workbook
    # ========================================================

    wb = load_workbook(output_file)

    ws = wb.active

    excel_header_row = header_row + 1

    # ========================================================
    # Tìm Address column trong sheet
    # ========================================================

    address_excel_col = None

    for cell in ws[excel_header_row]:
        if normalize_text(cell.value) == normalize_text(address_col):
            address_excel_col = cell.column

            break

    if not address_excel_col:
        raise ValueError("Không tìm thấy Address trong Excel gốc.")

    # ========================================================
    # Cột ngay sau Address
    # ========================================================

    maps_excel_col = address_excel_col + 1

    existing_header = safe_text(
        ws.cell(
            row=excel_header_row,
            column=maps_excel_col,
        ).value
    )

    # Nếu cột bên phải Address đang có data/header khác
    # thì insert một cột mới để không đè dữ liệu
    if existing_header and normalize_text(existing_header) not in {
        normalize_text("Google Maps"),
        normalize_text("google_maps_url"),
    }:
        ws.insert_cols(
            maps_excel_col,
            amount=1,
        )

    # ========================================================
    # Header
    # ========================================================

    header_cell = ws.cell(
        row=excel_header_row,
        column=maps_excel_col,
    )

    header_cell.value = "Google Maps"

    # Copy style từ Address
    address_header = ws.cell(
        row=excel_header_row,
        column=address_excel_col,
    )

    try:
        from copy import copy

        header_cell.font = copy(address_header.font)

        header_cell.fill = copy(address_header.fill)

        header_cell.border = copy(address_header.border)

        header_cell.alignment = copy(address_header.alignment)

        header_cell.number_format = address_header.number_format

    except Exception:
        pass

    # ========================================================
    # Write Address + Maps
    # ========================================================

    data_start_row = excel_header_row + 1

    for df_index, row in df.iterrows():
        excel_row = data_start_row + df_index

        # Address
        ws.cell(
            row=excel_row,
            column=address_excel_col,
        ).value = safe_text(row.get(address_col))

        # Maps
        maps_url = safe_text(row.get(maps_col))

        map_cell = ws.cell(
            row=excel_row,
            column=maps_excel_col,
        )

        map_cell.value = maps_url

        if is_google_maps_url(maps_url):
            map_cell.hyperlink = maps_url

            map_cell.style = "Hyperlink"

    # ========================================================
    # Width / wrap
    # ========================================================

    from openpyxl.utils import (
        get_column_letter,
    )

    address_letter = get_column_letter(address_excel_col)

    maps_letter = get_column_letter(maps_excel_col)

    ws.column_dimensions[address_letter].width = 45

    ws.column_dimensions[maps_letter].width = 55

    for row_num in range(
        data_start_row,
        ws.max_row + 1,
    ):
        ws.cell(
            row=row_num,
            column=address_excel_col,
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        ws.cell(
            row=row_num,
            column=maps_excel_col,
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    # ========================================================
    # Save
    # ========================================================

    wb.save(output_file)


# ============================================================
# MAIN
# ============================================================


def main():

    print()
    print("=" * 70)
    print("GOOGLE MAPS -> CRM EXCEL MERGER")
    print("=" * 70)
    print()

    # ========================================================
    # Check files
    # ========================================================

    if not os.path.exists(MAIN_FILE):
        print(f"❌ Không thấy file: {MAIN_FILE}")

        return

    if not os.path.exists(MAPS_FILE):
        print(f"❌ Không thấy file: {MAPS_FILE}")

        return

    # ========================================================
    # Load main
    # ========================================================

    print(f"📘 CRM: {MAIN_FILE}")

    main_df, header_row = load_main_excel(MAIN_FILE)

    print(f"   Header row: {header_row + 1}")

    print(f"   Data rows: {len(main_df)}")

    print()

    # ========================================================
    # Load maps
    # ========================================================

    print(f"🗺 Maps: {MAPS_FILE}")

    (
        maps_df,
        maps_title_col,
        maps_google_col,
        maps_url_col,
    ) = load_maps_excel(MAPS_FILE)

    print(f"   Title column: {maps_title_col}")

    print(f"   Maps column: {maps_google_col}")

    print(f"   URL fallback: {maps_url_col}")

    print()

    # ========================================================
    # Lookup
    # ========================================================

    lookup = build_google_maps_lookup(
        maps_df,
        maps_title_col,
        maps_google_col,
        maps_url_col,
    )

    print(f"🔎 Valid Maps lookup: {len(lookup)}")

    print()

    # ========================================================
    # Merge
    # ========================================================

    result = merge_google_maps(
        main_df,
        lookup,
    )

    print("=" * 70)

    print(f"✅ GẮN GOOGLE MAPS: {result['matched']}")

    print(f"⏭ ĐÃ CÓ MAPS: {result['already_has']}")

    print(f"❌ KHÔNG TÌM THẤY MAPS: {result['not_found']}")

    print(f"⚠️ KHÔNG CÓ TITLE: {result['empty_title']}")

    print("=" * 70)

    print()

    # ========================================================
    # Save
    # ========================================================

    save_excel(
        original_file=MAIN_FILE,
        output_file=OUTPUT_FILE,
        df=result["df"],
        header_row=header_row,
        address_col=result["address_col"],
        maps_col=result["maps_col"],
    )

    print(f"💾 Output:")

    print(f"   {OUTPUT_FILE}")

    print()

    print("🎉 Hoàn tất.")


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    try:
        main()

    except KeyboardInterrupt:
        print()
        print("⛔ Đã dừng bởi người dùng.")

    except Exception as error:
        print()
        print("=" * 70)

        print("❌ ERROR:")

        print(repr(error))

        print("=" * 70)

        raise
